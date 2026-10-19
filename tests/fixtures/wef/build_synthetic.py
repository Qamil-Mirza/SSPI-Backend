"""Write the synthetic WEF workbook in this directory:

    WEF-GCIHH_synthetic.xlsx   (WEF_QUELEC, and through it AQELEC)

Entirely invented, not WEF data: only the layout of the WEF Global
Competitiveness Index Historical Dataset workbook is reproduced (sheet
names, the column header of the first sheet, ``Value`` and ``Rank`` rows,
``-`` placeholders, text year headers, integer and decimal cells, blanks).
Every value comes from the simple formulas below. The two metadata sheets
and the workbook properties say it is synthetic.

The original workbook is not distributed with the repository while its
redistribution terms are open (docs/data-governance-review.md, GOV-2); see
README.md.

``build(variant)`` also writes four invalid variants, never committed: the
golden generator and the adapter tests build them in memory to record what
the legacy cleaner and the new adapter do with a repeated row, a code that
is not three capital letters, a code with a digit, and a text value.

Run from the repository root; the output is byte-for-byte deterministic
for a given Python and zlib (fixed workbook properties, fixed zip entry
times); ``cells`` compares content across environments:

    .venv/bin/python tests/fixtures/wef/build_synthetic.py

tests/golden/test_golden_wef.py checks that the committed file is exactly
what this script writes.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).parent
FILE_NAME = "WEF-GCIHH_synthetic.xlsx"
NOTE = "SYNTHETIC TEST WORKBOOK - invented for SSPI parity tests; not WEF data"
HEADER = ("Economy ISO3", "Economy Name", "Indicator ID", "Indicator", "Attribute 1", "Attribute 2", "Attribute 3", "Partner", *(str(y) for y in range(2007, 2018)))
QUELEC = ("WEF.GCIHH.EOSQ064", "Quality of electricity supply, 1-7 (best)")
ROADS = ("WEF.GCIHH.EOSQ065", "Quality of road infrastructure, 1-7 (best)")  # another indicator: filtered out
VARIANTS = ("duplicate_row", "lowercase_code", "code_with_digit", "text_value")
FIXED_TIME = datetime(2000, 1, 1)


def rows(variant: str | None = None) -> list[tuple]:
    """The first sheet's data rows, in order."""

    def years(values):
        out = [None] * 11
        for year, value in values.items():
            out[year - 2007] = value
        return out

    def series(start, step, first=2007, last=2017, skip=()):
        return {y: round(start + step * (y - first), 2) for y in range(first, last + 1) if y not in skip}

    economies = [
        # code, name, values of WEF.GCIHH.EOSQ064 (Attribute 1 = Value)
        ("MYS", "Malaysia", {**series(5.1, 0.07, skip=(2009,)), 2010: 5.123456789012345, 2013: 4}),  # a gap, full precision, an integer cell
        ("AUT", "Austria", series(6.5, 0.1, first=2015)),  # starts late
        ("USA", "United States", {**series(5.8, 0.1), 2017: 7}),  # ends at the upper goalpost, an integer cell
        ("TCD", "Chad", {**series(1.4, 0.15, last=2012), 2007: 1}),  # starts at the lower goalpost, ends early
        ("QMA", "Synthetic economy A", series(3.0, 0.2, first=2010, last=2014)),  # not ISO 3166: no World Bank access data
        ("QRA", "Synthetic region", series(4.0, 0.05)),  # an aggregate-like code, kept as legacy kept it
        ("QMB", "Synthetic economy B", {}),  # every year blank
    ]
    out: list[tuple] = []
    for code, name, values in economies:
        if variant == "lowercase_code" and code == "QMB":
            code, values = "qmb", {2010: 3.5}
        if variant == "code_with_digit" and code == "QMB":
            code, values = "QM1", {2010: 3.5}
        if variant == "text_value" and code == "QMB":
            values = {2010: "x"}
        out.append((code, name, *QUELEC, "Value", "-", "-", "-", *years(values)))
        out.append((code, name, *QUELEC, "Rank", "-", "-", "-", *years({y: 10 + (y % 7) for y in values})))
        out.append((code, name, *ROADS, "Value", "-", "-", "-", *years({y: 2.5 for y in range(2007, 2018)})))
    if variant == "duplicate_row":
        out.append(("MYS", "Malaysia", *QUELEC, "Value", "-", "-", "-", *years({2011: 5.5})))
    return out


def build(variant: str | None = None) -> bytes:
    """The workbook's bytes (``variant`` one of :data:`VARIANTS`, or ``None`` for the committed file)."""
    from openpyxl import Workbook

    assert variant is None or variant in VARIANTS, variant
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(HEADER)
    for row in rows(variant):
        data.append(row)
    indicators = workbook.create_sheet("Indicator Metadata")
    indicators.append(("Indicator ID", "Indicator Name", "Description", "Attribute 1 Description", "Economies", "Time Range", "Last Updated"))
    for code, name in (QUELEC, ROADS):
        indicators.append((code, name, f"{name} ({NOTE})", "Value Type", 7, "2007 - 2017", "2000-01-01"))
    dataset = workbook.create_sheet("Dataset Metadata")
    dataset.append(("Dataset ID", "Title", "Dataset Provider", "Description", "Legal Text"))
    dataset.append(("SYNTHETIC", NOTE, "SSPI test suite", NOTE, "Invented test data; no third-party rights."))
    workbook.properties.creator = "SSPI test suite"
    workbook.properties.title = NOTE
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    raw = io.BytesIO()
    workbook.save(raw)
    # openpyxl stamps the zip entries and the modification property with the current time: rewrite both
    fixed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                content = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>", rb"<dcterms:modified\1>2000-01-01T00:00:00Z</dcterms:modified>", content)
            target.writestr(zipfile.ZipInfo(item.filename, date_time=FIXED_TIME.timetuple()[:6]), content, compress_type=zipfile.ZIP_DEFLATED)
    return fixed.getvalue()


def cells(content: bytes) -> dict[str, list[tuple]]:
    """Every sheet's cell values: equal for two workbooks with the same content, whatever zlib compressed them."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True)
    try:
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def main() -> None:
    (HERE / FILE_NAME).write_bytes(build())


if __name__ == "__main__":
    main()
