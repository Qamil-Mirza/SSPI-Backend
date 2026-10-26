"""Write the synthetic Fragile States Index workbooks in this directory:

    fsi-2018_synthetic.xlsx, fsi-2021_synthetic.xlsx, fsi-2023_synthetic.xlsx   (FSI_SECAPP, and through it SECAPP)

Entirely invented, not Fund for Peace data: only the layout of the annual
FSI workbooks is reproduced, in its three variants (a sheet named after the
year, ``Year`` as a 1 January date, the indicators in C-E-P-S-X order, a
"Change from Previous Year" column, as in 2018-2020; ``Year`` as an
integer, as in 2021; a sheet named ``Sheet1`` with the S indicators first,
as in 2023). Every score is made up. The country names are real, in the
spellings the FSI uses, because the country matching is part of what is
tested; the scores attached to them mean nothing.

The original workbooks are not distributed with the repository while their
terms of use are unreviewed (docs/data-governance-review.md, GOV-5); see
README.md.

``build(edition, variant)`` also writes invalid variants, never committed
(see :data:`VARIANTS`).

Run from the repository root; the output is byte-for-byte deterministic
for a given Python and zlib (fixed workbook properties, fixed zip entry
times):

    .venv/bin/python tests/fixtures/fsi/build_synthetic.py

tests/golden/test_golden_secapp.py checks that the committed files are
exactly what this script writes.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).parent
EDITIONS = (2018, 2021, 2023)
NOTE = "SYNTHETIC TEST WORKBOOK - invented for SSPI parity tests; not Fund for Peace data"
INDICATORS = ("C1: Security Apparatus", "C2: Factionalized Elites", "C3: Group Grievance", "E1: Economy", "E2: Economic Inequality", "E3: Human Flight and Brain Drain",
              "P1: State Legitimacy", "P2: Public Services", "P3: Human Rights", "S1: Demographic Pressures", "S2: Refugees and IDPs", "X1: External Intervention")
VARIANTS = {
    "unapproved_name": "a 2021 row for \"Atlantis\", a name neither the legacy matcher nor the approved corrections place",
    "duplicate_country": "a second 2021 row for Malaysia",
    "blank_value": "a 2021 row whose C1 cell is empty",
    "year_not_edition": "a 2021 row whose Year is 2020",
}
FIXED_TIME = datetime(2000, 1, 1)

# FSI spelling per edition (None: not in that edition), C1 per edition
COUNTRIES = (
    ({2018: "Malaysia", 2021: "Malaysia", 2023: "Malaysia"}, {2018: 6.25, 2021: 6.1, 2023: 5.9}),
    ({2018: "Austria", 2021: "Austria", 2023: "Austria"}, {2018: 1.5, 2021: 1.4, 2023: 1.4}),
    ({2018: "United States", 2021: "United States", 2023: "United States"}, {2018: 3.0, 2021: 3.6, 2023: 3.3}),
    ({2018: "Somalia", 2021: "Somalia", 2023: "Somalia"}, {2018: 10.0, 2021: 9.932962647670891, 2023: 10.0}),  # the lower goalpost; full precision
    ({2018: "Finland", 2021: "Finland", 2023: "Finland"}, {2018: 0.3, 2021: 0.5, 2023: 0.3}),
    ({2018: "North Korea", 2021: "North Korea", 2023: "North Korea"}, {2018: 8.0, 2021: 8.1, 2023: 8.2}),  # legacy matcher: KOR
    ({2018: "South Korea", 2021: "South Korea", 2023: "South Korea"}, {2018: 2.0, 2021: 1.9, 2023: 1.8}),
    ({2018: "Cote d'Ivoire", 2021: "Côte d'Ivoire", 2023: "Cote d'Ivoire"}, {2018: 7.0, 2021: 6.9, 2023: 6.8}),  # the accented spelling matches
    ({2018: "Macedonia", 2021: "North Macedonia", 2023: "Macedonia"}, {2018: 4.5, 2021: 4.4, 2023: 4.3}),
    ({2018: "Micronesia", 2021: "Micronesia", 2023: "Micronesia"}, {2018: 2.5, 2021: 2.6, 2023: 2.7}),
    ({2018: None, 2021: "Palestine", 2023: "Palestine"}, {2021: 7.5, 2023: 7.7}),  # listed from 2021
    ({2018: "Israel and West Bank", 2021: "Israel and West Bank", 2023: "Israel and West Bank"}, {2018: 5.5, 2021: 5.6, 2023: 5.8}),  # legacy rule: ISR
    ({2018: "Congo Democratic Republic", 2021: "Congo Democratic Republic", 2023: "Congo Democratic Republic"}, {2018: 9.0, 2021: 9.1, 2023: 9.2}),
    ({2018: "Congo Republic", 2021: "Congo Republic", 2023: "Congo Republic"}, {2018: 6.0, 2021: 6.2, 2023: 6.1}),
    ({2018: "Guinea Bissau", 2021: "Guinea Bissau", 2023: "Guinea Bissau"}, {2018: 7.25, 2021: 7.0, 2023: 7.1}),
    ({2018: "Kyrgyz Republic", 2021: "Kyrgyz Republic", 2023: "Kyrgyz Republic"}, {2018: 5.0, 2021: 5.0, 2023: 5.2}),
    ({2018: "Sri Lanka ", 2021: "Sri Lanka", 2023: "Sri Lanka"}, {2018: 6.75, 2021: 6.4, 2023: 6.5}),  # padded in 2018
    ({2018: "Serbia  ", 2021: "Serbia", 2023: "Serbia"}, {2018: 4.0, 2021: 4.1, 2023: 4.2}),  # two spaces in 2018
    ({2018: "Kosovo", 2021: "Kosovo", 2023: "Kosovo"}, {2018: 5.75, 2021: 5.6, 2023: 5.5}),  # legacy rule: XKX
)


def rows(edition: int, variant: str | None = None) -> list[tuple]:
    """(Country, C1) rows of one edition, in order."""
    out = [(names[edition], c1[edition]) for names, c1 in COUNTRIES if names[edition] is not None]
    if edition == 2021 and variant == "unapproved_name":
        out.append(("Atlantis", 5.0))
    if edition == 2021 and variant == "duplicate_country":
        out.append(("Malaysia", 6.2))
    if edition == 2021 and variant == "blank_value":
        out.append(("Mongolia", None))
    return out


def build(edition: int, variant: str | None = None) -> bytes:
    """One edition's workbook bytes (``variant`` one of :data:`VARIANTS`, or ``None`` for the committed file)."""
    from openpyxl import Workbook

    assert edition in EDITIONS and (variant is None or variant in VARIANTS), (edition, variant)
    order = INDICATORS[9:11] + INDICATORS[:9] + INDICATORS[11:] if edition == 2023 else INDICATORS
    header = ("Country", "Year", "Rank", "Total", *order) + (("Change from Previous Year",) if edition == 2018 else ())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1" if edition == 2023 else str(edition)
    sheet.append(header)
    for rank, (name, c1) in enumerate(rows(edition, variant), start=1):
        year = datetime(edition, 1, 1) if edition == 2018 else edition
        if variant == "year_not_edition" and name == "Austria":
            year = edition - 1
        others = {column: round(1 + ((rank * 7 + i * 3) % 90) / 10, 1) for i, column in enumerate(INDICATORS[1:])}
        values = {INDICATORS[0]: c1, **others}
        total = round(sum(v for v in values.values() if v is not None), 1)
        sheet.append((name, year, f"{rank}th", total, *(values[c] for c in order)) + ((round(0.1 * (rank % 5) - 0.2, 1),) if edition == 2018 else ()))
    workbook.properties.creator = "SSPI test suite"
    workbook.properties.title = NOTE
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    raw = io.BytesIO()
    workbook.save(raw)
    fixed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                content = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>", rb"<dcterms:modified\1>2000-01-01T00:00:00Z</dcterms:modified>", content)
            target.writestr(zipfile.ZipInfo(item.filename, date_time=FIXED_TIME.timetuple()[:6]), content, compress_type=zipfile.ZIP_DEFLATED)
    return fixed.getvalue()


def file_name(edition: int) -> str:
    return f"fsi-{edition}_synthetic.xlsx"


def cells(content: bytes) -> dict[str, list[tuple]]:
    """Every sheet's cell values, to compare content across environments."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True)
    try:
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def main() -> None:
    for edition in EDITIONS:
        (HERE / file_name(edition)).write_bytes(build(edition))


if __name__ == "__main__":
    main()
