"""Write the synthetic SIPRI sources in this directory:

    armstransfers_synthetic.json   (SIPRI_ARMEXP, and through it ARMEXP)
    milex_synthetic.json           (the legacy military-expenditure API answer: characterization of the legacy cleaner)
    milex_workbook_synthetic.xlsx  (SIPRI_MILEXP's approved source, the revised workbook: MILEXP-4)

Entirely invented, not SIPRI data: only the layout of the two API answers
the legacy collectors stored is reproduced. ``armstransfers`` is
``{"result": "<csv text>"}`` with the notes lines, the ``Exports by,`` header
(36 year columns, then period total and percentage columns), one row per
supplier, a ``Total world export`` row and trailing blank lines; ``milex``
is ``{"Data": [sections]}`` whose "Share of GDP" section has title rows, the
header ``Country, Notes, 1990.0 ... 2024.0``, region rows, country rows of
fractions written as text, ``...`` for unavailable and ``xxx`` for not
independent. The workbook copies the revised workbook's "Share of GDP"
sheet: a title row, the legend, the header ``Country, Notes, 1988 ... 2025``
(whole-number years), region rows, country rows of fractions stored as
numbers, ``...`` and ``xxx``, and published zeros for countries without
military expenditure. Every value is made up; supplier and country names are
real because the name matching is part of what is tested. Each file says it
is synthetic.

The original answers are not distributed with the repository
(docs/data-governance-review.md, GOV-7); see README.md.

``build(name, variant)`` also writes invalid variants and probes, never
committed (see :data:`VARIANTS`).

Run from the repository root:

    .venv/bin/python tests/fixtures/sipri/build_synthetic.py

tests/golden/test_golden_armexp.py and test_golden_milexp.py check that the
committed files are exactly what this script writes.
"""

from __future__ import annotations

import json
from pathlib import Path

HERE = Path(__file__).parent
FILES = {"armstransfers": "armstransfers_synthetic.json", "milex": "milex_synthetic.json", "milex_workbook": "milex_workbook_synthetic.xlsx"}
NOTE = "SYNTHETIC: invented values in the SIPRI layout, not SIPRI data"
VARIANTS = {
    "armstransfers": {
        "no_table": "the notes without the 'Exports by,' table",
        "duplicate_country": "a 'Republic of Korea' row beside 'South Korea': both are KOR",
        "probe_are_exports": "a United Arab Emirates row: a hard-coded reference-class country with observed exports",
    },
    "milex": {
        "probe_empty_share": "a 'Share of GDP' section with no country row",
        "north_korea_values": "'Korea, North' with published values: the legacy matcher gives it KOR, as 'Korea, South'",
        "probe_iceland_values": "an Icelandic value: the hard-coded Iceland zero then collides with an observation",
        "short_row": "a country row one cell shorter than the header",
    },
    "milex_workbook": {
        "iceland_published_zero": "Iceland published as 0 in every year, as in the revised workbook: the hard-coded Iceland zero then meets observations",
        "no_share_sheet": "no 'Share of GDP' sheet",
        "wrong_title": "a 'Share of GDP' sheet with another title",
        "two_header_rows": "a second row starting 'Country'",
        "header_not_consecutive": "a header that skips a year",
        "window_not_covered": "a header starting in 1995",
        "duplicate_name": "a second 'Austria' row",
        "duplicate_identity": "'Republic of Korea' beside 'Korea, South': both are KOR",
        "unknown_marker": "a cell 'n.a.', not one of SIPRI's markers",
        "blank_cell": "an empty cell in a country row",
        "unknown_name": "a country name with no identity",
        "region_with_values": "a region row holding a number",
        "text_value": "a number written as text",
    },
}
ARM_YEARS = list(range(1990, 2026))


def _arms_row(name, values):
    """``values``: {year: cell text}; the period total and percentage columns are invented filler."""
    cells = [values.get(year, "") for year in ARM_YEARS]
    total = sum(float(c) for c in cells if c.strip())
    return ",".join([name, *cells, f"{total:g}", "1%", f"{total:g}", "1%"])


ARMS = (
    ("United States", {y: str(10000 + 37 * (y - 1990)) for y in ARM_YEARS}),  # every year, far above the 500 goalpost
    ("Russia", {y: str(3000 + 11 * (y - 1992)) for y in range(1992, 2026)}),  # starts 1992
    ("Austria", {1990: "30", 1991: "20", 1995: "", 1999: "23", 2000: "23", 2004: "3", 2010: "75", 2020: "10", 2023: "4"}),  # gaps
    ("Malaysia", {1995: "0 ", 1999: "11", 2013: "15", 2014: "15", 2019: "1", 2024: "20"}),  # a published zero written "0 "
    ("South Korea", {y: str(100 + y - 1990) for y in range(1995, 2026)}),
    ("North Korea", {1990: "6", 1991: "85", 2010: "7", 2014: "1", 2023: "9", 2024: "211"}),  # legacy rule: PRK
    ("Belarus", {1995: "400", 1996: "500", 1997: "600", 1998: "250"}),  # history before 2000 only: carried forward 1999-2023
    ("Brazil", {2024: "120", 2025: "130"}),  # after 2023 only: carried back 2000-2023 from 2024
    ("Viet Nam", {2005: "499.5", 2006: "500", 2007: "1000.25"}),  # at and around the goalpost
    ("Soviet Union", {1990: "9000", 1991: "8000"}),  # skip list
    ("Brunei", {2015: "2"}),  # the legacy matcher does not place these two; the approved SIPRI_ARMEXP aliases do (BRN, BIH)
    ("Bosnia-Herzegovina", {2003: "9", 2011: "0 ", 2012: "14"}),
    ("European Union**", {2020: "0 "}),  # skip list
    ("African Union**", {2023: "1"}),  # the legacy matcher does not place it
)


def armstransfers_text(variant=None):
    lines = [
        NOTE,
        "Volume of transfers of major arms",
        "Figures are in millions of SIPRI trend-indicator values (TIVs).",
        "A '0' indicates that the volume of deliveries is between 0 and 0.5 million SIPRI TIV. An empty field indicates that no deliveries have been identified.",
        "",
        "Source: SIPRI Arms Transfers Database (c) SIPRI.",
        "Data generated: 26 Oct 2026 8:50:40 PM",
        "",
    ]
    if variant != "no_table":
        lines.append(",".join(["Exports by", *map(str, ARM_YEARS), "1990-2025", "Percentage", "Sum total years", "Percentage of total"]))
        rows = list(ARMS)
        if variant == "duplicate_country":
            rows.append(("Republic of Korea", {2010: "50"}))
        if variant == "probe_are_exports":
            rows.append(("United Arab Emirates", {2001: "28", 2005: "6", 2015: "143"}))
        lines.extend(_arms_row(name, values) for name, values in rows)
        lines.append(_arms_row("Total world export", {y: "40000" for y in ARM_YEARS}))
    return "\n".join(lines) + "\n\n"


MILEX_YEARS = [f"{y}.0" for y in range(1990, 2025)]


def _share(first, step, *, gaps=(), xxx=()):
    out = []
    for i in range(35):
        year = 1990 + i
        out.append("xxx" if year in xxx else ". ." if year in gaps else "..." if year < first else repr(round(step * (1 + (i % 7) / 10) + 0.001 * i, 6)))
    return out


SHARE = (
    ("Algeria", "§", _share(1990, 0.03)),  # a recognised note mark: values start at column 2
    ("Libya", "‡§¶", _share(1997, 0.02)),  # three marks
    ("Austria", "", _share(1990, 0.009)),  # empty notes: the legacy cleaner reads the empty cell as a value
    ("Malaysia", "", _share(1990, 0.015, gaps=(2001, 2002))),
    ("United States of America", "", _share(1990, 0.035)),
    ("Korea, North", "", ["..."] * 35),
    ("Korea, South", "", _share(1990, 0.025)),
    ("Ukraine", "†", _share(1992, 0.02, xxx=(1990, 1991))),  # a mark the legacy cleaner does not recognise
    ("Iceland", "", ["..."] * 35),  # no military
    ("Costa Rica", "", ["..."] * 35),
    ("Eritrea", "", ["0.21", "1.5", *_share(1992, 0.2)[2:]]),  # a value of 1.5: not below 1, kept as written
    ("Brunei", "§", _share(1990, 0.03)),  # the legacy matcher does not place these four names
    ("Congo, Republic", "§", _share(1992, 0.02)),
    ("Cote d'Ivoire", "", _share(1990, 0.012)),
    ("Gambia, The", "‡§", _share(1990, 0.005)),
    ("German Democratic Republic", "", ["xxx"] * 35),  # the legacy matcher's "democratic republic" rule: COD
    ("Congo, DR", "", _share(1990, 0.011)),
)


def milex_document(variant=None):
    share_rows = [
        [NOTE],
        ["Military expenditure by country as percentage of gross domestic product, 1990-2024     © SIPRI 2025"],
        ["Countries are grouped by region and subregion"],
        ["\". .\" = data unavailable. \"xxx\" = country did not exist or was not independent during all or part of the year in question."],
        [""],
        ["Country", "Notes", *MILEX_YEARS],
        ["Africa"],
        ["North Africa"],
    ]
    if variant != "probe_empty_share":
        for name, notes, values in SHARE:
            if variant == "north_korea_values" and name == "Korea, North":
                values = _share(2005, 0.2)
            if variant == "probe_iceland_values" and name == "Iceland":
                values = ["..."] * 30 + ["0.001", "...", "...", "...", "..."]
            row = [name, notes, *values]
            if variant == "short_row" and name == "Austria":
                row = row[:-1]
            share_rows.append(row)
            if name == "Libya":
                share_rows.extend([["Europe"], ["Western Europe"]])
    return {
        "Data": [
            {"DataType": "Front page", "Rows": [[]]},
            {"DataType": "Local currency calendar years", "Rows": [[NOTE], ["Country", "Currency", "Notes", *MILEX_YEARS], ["Austria", "Euro", "", *["1.0E9"] * 35]]},
            {"DataType": "Share of GDP", "Rows": share_rows},
            {"DataType": "Footnotes", "Rows": [["Footnotes and special notes"], ["§", "Figures for these countries are for the adopted budget, rather than actual expenditure"]]},
        ]
    }


WORKBOOK_YEARS = list(range(1988, 2026))  # wider than the 1990-2024 window read
WORKBOOK_TITLE = "Military expenditure by country as percentage of gross domestic product, 1988–2025  © SIPRI 2026"


def _fractions(first, step, *, gaps=(), xxx=(), zero_from=None):
    out = []
    for i, year in enumerate(WORKBOOK_YEARS):
        if year in xxx:
            out.append("xxx")
        elif year in gaps or year < first:
            out.append("...")
        elif zero_from is not None and year >= zero_from:
            out.append(0)
        else:
            out.append(round(step * (1 + (i % 5) / 10) + 0.0001 * i, 9))
    return out


WORKBOOK = (
    ("Algeria", "§", _fractions(1988, 0.03)),
    ("Libya", "‡§¶", _fractions(1997, 0.02, gaps=(2014, 2015, 2016))),
    ("Austria", "", _fractions(1988, 0.009)),
    ("Malaysia", "", _fractions(1988, 0.015, gaps=(2001, 2002))),
    ("United States of America", "44", _fractions(1988, 0.035)),
    ("Korea, North", "", ["..."] * len(WORKBOOK_YEARS)),
    ("Korea, South", "", _fractions(1988, 0.025)),
    ("Ukraine", "†", _fractions(1992, 0.02, xxx=(1988, 1989, 1990, 1991))),
    ("Belarus", "", _fractions(1992, 0.012, xxx=(1988, 1989, 1990, 1991))),
    ("Iceland", "86", ["..."] * len(WORKBOOK_YEARS)),  # the committed file keeps Iceland collision-free; see the variant
    ("Costa Rica", "37", [0] * len(WORKBOOK_YEARS)),  # a published zero: no military expenditure (footnote)
    ("Panama", "43", _fractions(1988, 0.01, zero_from=2000)),  # values, then published zeros
    ("Eritrea", "", [*_fractions(1993, 0.2)[:3], 1.5, *_fractions(1993, 0.2)[4:]]),  # 1991: a share of 1 or more (150%)
    ("Venezuela", "‖", _fractions(1988, 0.015, gaps=(2018, 2019, 2020, 2021, 2022, 2023))),
    ("Brunei", "§", _fractions(1988, 0.03)),  # the approved name table places these four (MILEXP-1)
    ("Congo, Republic", "§", _fractions(1992, 0.02)),
    ("Cote d'Ivoire", "", _fractions(1988, 0.012)),
    ("Gambia, The", "‡§", _fractions(1988, 0.005)),
    ("Kosovo", "", _fractions(2008, 0.008, xxx=tuple(range(1988, 2008)))),  # the legacy matcher's XKX
    ("German Democratic Republic", "", [0.05, 0.048, *["xxx"] * (len(WORKBOOK_YEARS) - 2)]),  # excluded, never DEU
    ("USSR", "", [0.12, 0.11, 0.1, 0.09, *["xxx"] * (len(WORKBOOK_YEARS) - 4)]),  # excluded
    ("Congo, DR", "", _fractions(1988, 0.011)),
)


def workbook_rows(variant=None):
    header = ["Country", "Notes", *WORKBOOK_YEARS]
    if variant == "header_not_consecutive":
        header = ["Country", "Notes", *[y for y in WORKBOOK_YEARS if y != 2005], 2026]
    if variant == "window_not_covered":
        header = ["Country", "Notes", *range(1995, 1995 + len(WORKBOOK_YEARS))]
    width = len(header)
    rows = [
        [WORKBOOK_TITLE if variant != "wrong_title" else "Military expenditure by country in constant (2024) US$ m."],
        [NOTE],
        ["Figures in blue are SIPRI estimates. Figures in red indicate highly uncertain data."],
        ["\". .\" = data unavailable. \"xxx\" = country did not exist or was not independent during all or part of the year in question."],
        [""],
        header,
        [None] * width,
        ["Africa", *[None] * (width - 1)],
        ["North Africa", *[None] * (width - 1)],
    ]
    for name, notes, values in WORKBOOK:
        values = list(values)
        if variant == "iceland_published_zero" and name == "Iceland":
            values = [0] * len(WORKBOOK_YEARS)
        if variant == "unknown_marker" and name == "Austria":
            values[WORKBOOK_YEARS.index(2005)] = "n.a."
        if variant == "blank_cell" and name == "Austria":
            values[WORKBOOK_YEARS.index(2005)] = None
        if variant == "text_value" and name == "Austria":
            values[WORKBOOK_YEARS.index(2005)] = "0.009"
        rows.append([name, notes, *values])
        if name == "Libya":
            rows.extend([["Europe", *[None] * (width - 1)], ["East Asia", None, *["..."] * len(WORKBOOK_YEARS)]])
    if variant == "region_with_values":
        rows.append(["Middle East", None, *[0.02] * len(WORKBOOK_YEARS)])
    if variant == "two_header_rows":
        rows.append(list(header))
    if variant == "duplicate_name":
        rows.append(["Austria", "", *_fractions(1988, 0.008)])
    if variant == "duplicate_identity":
        rows.append(["Republic of Korea", "", *_fractions(1988, 0.02)])
    if variant == "unknown_name":
        rows.append(["Atlantis", "", *_fractions(1988, 0.02)])
    return [row + [None] * (width - len(row)) for row in rows]


def workbook_bytes(variant=None) -> bytes:
    """The workbook's bytes, with fixed timestamps so the committed file is reproducible."""
    import datetime
    import io
    import re
    import zipfile

    from openpyxl import Workbook

    fixed_time = datetime.datetime(2000, 1, 1)  # noqa: DTZ001 - a naive constant, as the zip format stores it
    book = Workbook()
    front = book.active
    front.title = "Front page"
    front.append([NOTE])
    share = book.create_sheet("Share of GDP (synthetic)" if variant == "no_share_sheet" else "Share of GDP")
    for row in workbook_rows(variant):
        share.append(row)
    footnotes = book.create_sheet("Footnotes")
    footnotes.append([NOTE])
    footnotes.append([37, "Costa Rica: synthetic footnote: no armed forces, no military expenditure."])
    footnotes.append([43, "Panama: synthetic footnote: no military expenditure from 2000."])
    footnotes.append([86, "Iceland: synthetic footnote: no standing military, no military expenditure."])
    book.properties.creator = "SSPI test suite"
    book.properties.title = NOTE
    book.properties.created = fixed_time
    book.properties.modified = fixed_time
    raw = io.BytesIO()
    book.save(raw)
    fixed = io.BytesIO()  # openpyxl stamps the zip entries and the modification property with the current time: rewrite both
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                content = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>", rb"<dcterms:modified\1>2000-01-01T00:00:00Z</dcterms:modified>", content)
            target.writestr(zipfile.ZipInfo(item.filename, date_time=fixed_time.timetuple()[:6]), content, compress_type=zipfile.ZIP_DEFLATED)
    return fixed.getvalue()


def build(name, variant=None) -> bytes:
    if variant is not None and variant not in VARIANTS[name]:
        raise KeyError(variant)
    if name == "milex_workbook":
        return workbook_bytes(variant)
    document = {"result": armstransfers_text(variant)} if name == "armstransfers" else milex_document(variant)
    return (json.dumps(document, ensure_ascii=False, indent=1) + "\n").encode("utf-8")


if __name__ == "__main__":
    for name, file_name in FILES.items():
        (HERE / file_name).write_bytes(build(name))
        print(f"wrote {HERE / file_name}")
